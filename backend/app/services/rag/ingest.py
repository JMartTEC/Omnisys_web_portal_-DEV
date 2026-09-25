"""Extracción de texto y fragmentación (chunking) de documentos cargados."""
import csv
import io
import json

SOPORTADOS = {"pdf", "docx", "xlsx", "csv", "md", "txt", "json"}


def extraer_texto(nombre: str, contenido: bytes) -> str:
    ext = nombre.rsplit(".", 1)[-1].lower()
    if ext not in SOPORTADOS:
        raise ValueError(f"Formato .{ext} no soportado. Usa: {', '.join(sorted(SOPORTADOS))}")
    if ext in {"txt", "md"}:
        return contenido.decode("utf-8", errors="ignore")
    if ext == "json":
        return json.dumps(json.loads(contenido.decode("utf-8")), ensure_ascii=False, indent=1)
    if ext == "csv":
        filas = csv.reader(io.StringIO(contenido.decode("utf-8-sig", errors="ignore")))
        return "\n".join(" | ".join(f) for f in filas)
    if ext == "pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(contenido)).pages)
    if ext == "docx":
        from docx import Document
        d = Document(io.BytesIO(contenido))
        partes = [p.text for p in d.paragraphs if p.text and p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                partes.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(partes)
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        partes = []
        for ws in wb.worksheets:
            partes.append(f"## Hoja {ws.title}")
            encabezado = None
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v).strip() for v in row]
                if not any(vals):
                    continue
                if encabezado is None:
                    encabezado = vals
                    continue
                # "columna: valor" conserva semántica de diccionarios y glosarios
                partes.append("; ".join(f"{h}: {v}" for h, v in zip(encabezado, vals) if v and h))
        return "\n".join(partes)
    return ""


def fragmentar(texto: str, tamano: int = 900, traslape: int = 150) -> list[str]:
    """Divide por párrafos/líneas respetando un tamaño máximo con traslape."""
    bloques = [b.strip() for b in texto.replace("\r", "").split("\n") if b.strip()]
    chunks, actual = [], ""
    for b in bloques:
        if len(actual) + len(b) + 1 <= tamano:
            actual = f"{actual}\n{b}" if actual else b
            continue
        if actual:
            chunks.append(actual)
            actual = actual[-traslape:] + "\n" + b if traslape else b
        else:
            actual = b
        while len(actual) > tamano:
            chunks.append(actual[:tamano])
            actual = actual[tamano - traslape:]
    if actual.strip():
        chunks.append(actual)
    return chunks
